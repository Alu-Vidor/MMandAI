"""Build applicant-facing data from the current curriculum and editorial content."""
from pathlib import Path
import argparse, json, re
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('--plan', type=Path)
args = parser.parse_args()
plans = [args.plan] if args.plan else list(Path('C:/Users/danii/Desktop/РПД/МатМодиИИ/Учебный план').glob('*.xlsx'))
if len(plans) != 1:
    raise ValueError('Provide the authoritative workbook with --plan')
plan = plans[0]
workbook = load_workbook(plan, data_only=True)
sheet = workbook['ПланСвод']
assert [sheet.cell(2, c).value for c in range(18, 28)] == [f'Семестр {s}' for s in [1,2,3,4,5,6,7,8,9,'A']]
content = json.loads((ROOT/'tools/course_content.json').read_text(encoding='utf-8'))
examples = json.loads((ROOT/'tools/practical_examples.json').read_text(encoding='utf-8'))
old_text = (ROOT/'data.js').read_text(encoding='utf-8')
old = json.JSONDecoder().raw_decode(old_text.split('window.COURSES = ', 1)[1])[0]
groups = json.JSONDecoder().raw_decode(old_text.split('window.GROUPS = ', 1)[1])[0]
courses, counted = [], []
for row in sheet.iter_rows(min_row=7):
    code, name = row[1].value, row[2].value
    if not isinstance(code, str) or not name:
        continue
    if not (re.fullmatch(r'Б1\.(?:О|В)\.(?:\d{2}\.\d{2}|ДВ\.\d{2}\.\d{2})', code)
            or code.startswith(('ФТД.', 'Б3.')) or (code.startswith('Б2.') and '(' in code)):
        continue
    semesters = [i+1 for i,c in enumerate(row[17:27]) if c.value is not None and c.value != 0]
    if not semesters:
        # The sports elective has hours but zero credits; use its assessment semesters.
        controls = ''.join(str(c.value or '') for c in row[3:8])
        semesters = sorted({10 if ch == 'A' else int(ch) for ch in controls if ch in '123456789A'})
    assert semesters and name in content and name in examples, (code, name)
    kind = ('Факультатив' if code.startswith('ФТД.') else
            'По выбору · пара '+('1' if '.ДВ.01.' in code else '2') if '.ДВ.' in code else
            'Итоговая работа' if code.startswith('Б3.') else
            'Практика' if code.startswith('Б2.') else
            'Обязательная дисциплина' if '.О.' in code else 'Профильная дисциплина')
    course = dict(id=code, name=name, semesters=semesters, credits=float(row[9].value or 0),
                  hours=int(row[11].value or 0), **content[name], example=examples[name], kind=kind)
    courses.append(course)
    if row[0].value == '+' and not code.startswith('ФТД.'):
        counted.append(course)
assert len(courses) == len({c['id'] for c in courses})
assert sum(c['credits'] for c in counted) == 300
assert sum(c['credits'] for c in counted if c['id'].startswith('Б2.')) == 27
old_by_name = {c['name']:c for c in old}
new_by_name = {c['name']:c for c in courses}
changes = {
    'source': plan.name,
    'added': sorted(new_by_name.keys()-old_by_name.keys()),
    'removed': sorted(old_by_name.keys()-new_by_name.keys()),
    'changed': {name:{key:{'before':old_by_name[name][key], 'after':new_by_name[name][key]}
                      for key in ('id','semesters','credits','hours') if old_by_name[name][key]!=new_by_name[name][key]}
                for name in new_by_name.keys() & old_by_name.keys()
                if any(old_by_name[name][key]!=new_by_name[name][key] for key in ('id','semesters','credits','hours'))}}
(ROOT/'source-notes').mkdir(exist_ok=True)
(ROOT/'source-notes/curriculum-diff.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
# Rebind cross-links by course name before replacing the IDs in generated data.
aliases = {'Программирование математических вычислений':'Программирование и прикладная математика',
           'Прикладные задачи высшей математики':'Компьютерные технологии в математике'}
html_path = ROOT/'index.html'
html = html_path.read_text(encoding='utf-8')
old_by_id = {c['id']:c for c in old}
def rebind(match):
    name = aliases.get(old_by_id[match[1]]['name'], old_by_id[match[1]]['name'])
    assert name in new_by_name, f'Removed linked course: {name}'
    return 'data-course="'+new_by_name[name]['id']+'"'
html_path.write_text(re.sub(r'data-course="([^"]+)"', rebind, html),encoding='utf-8')
(ROOT/'data.js').write_text('window.COURSES = '+json.dumps(courses,ensure_ascii=False,indent=2)+
                          ';\nwindow.GROUPS = '+json.dumps(groups,ensure_ascii=False)+';\n',encoding='utf-8')
print(json.dumps({'courses':len(courses),'credits':300,'added':changes['added'],
                  'removed':changes['removed'],'changed':len(changes['changed'])},ensure_ascii=False))
