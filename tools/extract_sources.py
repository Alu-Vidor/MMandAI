from pathlib import Path
import json, csv, zipfile, re
from docx import Document
from openpyxl import load_workbook
root=Path('C:/Users/danii/Desktop/РПД/МатМодиИИ')
out=Path('source-notes'); out.mkdir(exist_ok=True)
wb=load_workbook(next(root.glob('*.xlsx')),data_only=True)
with (out/'plan.txt').open('w',encoding='utf-8') as f:
 for ws in wb:
  f.write('\nSHEET '+ws.title+'\n')
  for row in ws:
   vals=['%s:%s'%(c.column_letter,c.value) for c in row if c.value is not None]
   if vals:f.write(str(row[0].row)+' '+' | '.join(vals)+'\n')
records=[]
for r in csv.DictReader((root/'РПД_2026_итоговые/Реестр_дисциплин.csv').open(encoding='utf-8-sig'),delimiter=';'):
 d=Document(root/'РПД_2026_итоговые'/r['Файл'])
 txt='\n'.join([p.text for p in d.paragraphs]+[' | '.join(c.text for c in row.cells) for t in d.tables for row in t.rows])
 (out/(r['Код']+'.txt')).write_text(txt,encoding='utf-8')
 records.append({'id':r['Код'],'name':r['Дисциплина'],'semesters':[int(s.strip()) for s in r['Семестры'].split(',')],'credits':int(r['Зачетные единицы']),'hours':int(r['Часы'])})
(out/'registry.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
print(wb.sheetnames)
print('Extracted',len(records),'programmes')
